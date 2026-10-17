import csv
import io
import json
import re
import traceback
import zipfile
from datetime import date, datetime
from xml.etree import ElementTree as ET

from django.db import transaction

from .models import Application

PREVIEW_ROWS = 5
SUPPORTED_EXTENSIONS = (".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".ods", ".json")

# Header names (normalized: lowercase, dots removed, whitespace collapsed) per field.
# The app's own German and English field labels are included, so an exported
# or template file round-trips.
FIELD_ALIASES = {
    "company": ["company", "unternehmen", "firma", "firmenname", "arbeitgeber", "employer",
                "unternehmen anschrift", "firma anschrift"],
    "position": ["position", "stelle", "job", "jobtitel", "job title", "stellenbezeichnung", "stellentitel"],
    "date": ["date", "datum", "beworben am", "bewerbungsdatum", "application date"],
    "contact_person": ["contact person", "contact", "ansprechpartner", "ansprechpartnerin",
                       "ansprechperson", "kontaktperson", "kontakt"],
    "email": ["email", "e-mail", "mail", "e mail"],
    "phone": ["phone", "telefon", "tel", "telefonnummer", "telephone"],
    "follow_up": ["follow up on", "follow up", "followup", "nachfassen", "nachfassen am"],
    "interview_date": ["interview on", "interview date", "gespräch am", "gespraech am", "vorstellungsgespräch"],
    "response_date": ["response received on", "response date", "antwort erhalten am", "antwort am", "antwort"],
    "source": ["found via", "source", "gefunden über", "gefunden ueber", "quelle", "fundstelle"],
    "channel": ["application channel", "channel", "bewerbungsweg", "bewerbungsart"],
    "work_mode": ["work model", "work mode", "arbeitsmodell", "arbeitsform"],
    "salary": ["salary", "gehalt", "gehaltsvorstellung"],
    "street": ["street and number", "street", "straße und hausnummer", "strasse und hausnummer",
               "straße", "strasse"],
    "postal_code": ["postal code", "zip", "plz", "postleitzahl"],
    "city": ["city", "ort", "stadt"],
    "status": ["status"],
    "link": ["job posting link", "link", "url", "website", "webseite", "link zur stelle", "stellenanzeige"],
    "rejection_reason": ["rejection reason", "absagegrund", "grund der absage"],
    "notes": ["notes", "notizen", "bemerkung", "bemerkungen", "kommentar", "kommentare"],
    "favorite": ["favorite", "favorit", "favoriten", "stern"],
}
ALIAS_LOOKUP = {alias: key for key, aliases in FIELD_ALIASES.items() for alias in aliases}
COMPANY_ALIASES = set(FIELD_ALIASES["company"])

DATE_FIELDS = ("date", "follow_up", "interview_date", "response_date")
TEXT_FIELDS = ("company", "position", "contact_person", "email", "phone", "link", "street",
               "postal_code", "city", "source", "salary", "rejection_reason", "notes")

STATUS_LABELS = {
    "planned": "planned", "geplant": "planned",
    "unsure": "unsure", "unsicher": "unsure",
    "applied": "applied", "beworben": "applied",
    "interview": "interview", "gespräch": "interview", "gespraech": "interview",
    "rejected": "rejected", "absage": "rejected",
    "offer": "offer", "zusage": "offer",
}
CHANNEL_LABELS = {
    "email": "email", "e-mail": "email",
    "portal": "portal",
    "post": "post", "mail": "post", "brief": "post",
    "phone": "phone", "telefon": "phone",
    "personal": "personal", "persönlich": "personal", "persoenlich": "personal", "in person": "personal",
}
WORK_MODE_LABELS = {
    "onsite": "onsite", "on-site": "onsite", "vor ort": "onsite", "präsenz": "onsite",
    "hybrid": "hybrid",
    "remote": "remote",
}

FAVORITE_VALUES = {"ja", "yes", "true", "wahr", "1", "x", "y", "j", "★"}

DATE_FORMATS = ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")
POSTAL_RE = re.compile(r"^(\d{4,5})\s+(.+)$")
SALUTATION_RE = re.compile(r"^(?:herr|frau)\s+(.+)$", re.IGNORECASE)


class ImportFileError(Exception):
    def __init__(self, code):
        super().__init__(code)
        self.code = code


def normalize(value):
    text = re.sub(r"[\s_]+", " ", str(value or "").strip().lower())
    return text.replace(".", "")


def as_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_date(value):
    if value is None or as_text(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = as_text(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError(text)


def clip(field, value):
    limit = Application._meta.get_field(field).max_length
    return value[:limit] if limit else value


def decode_text(raw):
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252")


def parse_delimited(text):
    try:
        delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ";"
    return list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))


def read_csv(upload):
    return parse_delimited(decode_text(upload.read())), None


def read_txt(upload):
    """A .txt is either a delimited table with a header row or one company name per line."""
    text = decode_text(upload.read())
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return [], None
    if any(normalize(cell) in ALIAS_LOOKUP for cell in re.split(r"[;\t,]", lines[0])):
        return parse_delimited(text), None
    return [["company"]] + [[line] for line in lines], None


def read_json(upload):
    data = json.loads(decode_text(upload.read()))
    if isinstance(data, dict):
        data = next((v for v in data.values() if isinstance(v, list)), [])
    if data and all(isinstance(item, str) for item in data):
        return [["company"]] + [[item] for item in data], None
    headers = []
    for item in data:
        if isinstance(item, dict):
            headers.extend(k for k in item if k not in headers)
    rows = [headers] + [[item.get(h) for h in headers] for item in data if isinstance(item, dict)]
    return rows, None


def read_xlsx(upload):
    from openpyxl import load_workbook

    sheet = load_workbook(upload, data_only=True).worksheets[0]
    rows, targets = [], []
    for row in sheet.iter_rows():
        rows.append([cell.value for cell in row])
        # A cell showing text like "Karriere" may hide the real URL in its hyperlink.
        targets.append([getattr(getattr(cell, "hyperlink", None), "target", None) for cell in row])
    return rows, targets


ODS_NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "xlink": "http://www.w3.org/1999/xlink",
}


def q(prefix, name):
    return f"{{{ODS_NS[prefix]}}}{name}"


def ods_text(cell):
    paragraphs = []
    for p in cell.findall(q("text", "p")):
        parts = []

        def walk(element):
            if element.text:
                parts.append(element.text)
            for child in element:
                if child.tag == q("text", "s"):
                    parts.append(" " * int(child.get(q("text", "c"), 1)))
                elif child.tag == q("text", "tab"):
                    parts.append("\t")
                elif child.tag == q("text", "line-break"):
                    parts.append("\n")
                else:
                    walk(child)
                if child.tail:
                    parts.append(child.tail)

        walk(p)
        paragraphs.append("".join(parts))
    return "\n".join(paragraphs)


def ods_value(cell):
    kind = cell.get(q("office", "value-type"))
    if kind in ("float", "percentage", "currency"):
        return float(cell.get(q("office", "value")))
    if kind == "date":
        raw = cell.get(q("office", "date-value"))
        try:
            return datetime.fromisoformat(raw)
        except ValueError:
            return raw
    return ods_text(cell) or None


def read_ods(upload):
    """Read the first sheet of an OpenDocument spreadsheet (LibreOffice Calc)."""
    with zipfile.ZipFile(upload) as archive:
        root = ET.fromstring(archive.read("content.xml"))
    table = root.find(f".//{q('table', 'table')}")
    cell_tags = (q("table", "table-cell"), q("table", "covered-table-cell"))
    rows, targets = [], []
    for row in table.iter(q("table", "table-row")):
        repeat = int(row.get(q("table", "number-rows-repeated"), 1))
        values, links = [], []
        for cell in row:
            if cell.tag not in cell_tags:
                continue
            count = min(int(cell.get(q("table", "number-columns-repeated"), 1)), 200)
            anchor = cell.find(f".//{q('text', 'a')}")
            values.extend([ods_value(cell)] * count)
            links.extend([anchor.get(q("xlink", "href")) if anchor is not None else None] * count)
        while values and values[-1] is None:
            values.pop()
            links.pop()
        if not values and repeat > 1000:
            break  # the long run of empty filler rows at the end of a sheet
        for _ in range(min(repeat, 1000)):
            rows.append(values)
            targets.append(links)
    return rows, targets


READERS = {
    ".csv": read_csv, ".tsv": read_csv, ".txt": read_txt, ".json": read_json,
    ".xlsx": read_xlsx, ".xlsm": read_xlsx, ".ods": read_ods,
}


def read_rows(upload):
    extension = "." + upload.name.lower().rsplit(".", 1)[-1]
    try:
        return READERS[extension](upload)
    except Exception as exc:
        traceback.print_exc()  # the real reason shows up in the Django terminal
        raise ImportFileError("unreadable") from exc


def company_key(name):
    """Companies count as the same if they match ignoring case and extra spaces."""
    return re.sub(r"\s+", " ", name).strip().lower()


def split_company_cell(text):
    """Split a multi-line cell like 'Company / Herr Name / Street 1 / 12345 City'."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    parts = {"company": lines[0]}
    leftover = []
    for line in lines[1:]:
        postal = POSTAL_RE.match(line)
        salutation = SALUTATION_RE.match(line)
        if postal and "postal_code" not in parts:
            parts["postal_code"] = postal.group(1)
            parts["city"] = postal.group(2).split(",")[0].strip()
        elif salutation and "contact_person" not in parts:
            parts["contact_person"] = salutation.group(1).strip()
        elif re.search(r"\d", line) and "street" not in parts:
            parts["street"] = line
        else:
            leftover.append(line)
    if leftover:
        parts["notes"] = "\n".join(leftover)
    return parts


def get(row, columns, key):
    idx = columns.get(key)
    return row[idx] if idx is not None and idx < len(row) else None


def run_import(upload, dry_run=False, default_status="planned", update_existing=True):
    rows, targets = read_rows(upload)
    # The header is the first row that names a company column. Rows above it (a title,
    # blank lines) are ignored. Without such a row, the first non-empty row is used.
    header_row = next(
        (i for i, r in enumerate(rows[:20]) if any(normalize(v) in COMPANY_ALIASES for v in r)),
        None,
    )
    if header_row is None:
        header_row = next((i for i, r in enumerate(rows) if any(as_text(v) for v in r)), None)
    if header_row is None:
        raise ImportFileError("empty")

    columns, ignored = {}, []
    for idx, header in enumerate(rows[header_row]):
        key = ALIAS_LOOKUP.get(normalize(header))
        if key:
            columns.setdefault(key, idx)
        elif as_text(header):
            ignored.append(as_text(header))
    if "company" not in columns:
        raise ImportFileError("no_company_column")

    # Existing applications by company. If a company appears more than once, the oldest entry is updated.
    by_company = {}
    for application in Application.objects.order_by("id"):
        by_company.setdefault(company_key(application.company), application)

    seen = set()  # companies already handled in this file
    to_create, to_update, errors, warnings, preview, update_preview = [], [], [], [], [], []
    duplicates = unchanged = skipped = 0

    for number, row in enumerate(rows[header_row + 1:], start=header_row + 2):
        if not any(as_text(v) for v in row):
            continue

        data = {key: as_text(get(row, columns, key)) for key in TEXT_FIELDS}

        if targets and "link" in columns:
            target_row = targets[number - 1]
            idx = columns["link"]
            if idx < len(target_row) and target_row[idx]:
                data["link"] = target_row[idx]

        if "\n" in data["company"]:
            parts = split_company_cell(data["company"])
            data["company"] = parts.pop("company")
            for key, value in parts.items():
                if not data.get(key):
                    data[key] = value

        if not data["company"]:
            errors.append({"row": number, "code": "missing_company"})
            continue

        data["email"] = re.sub(r"^mailto:", "", data["email"], flags=re.IGNORECASE)
        data["phone"] = re.sub(r"^tel:", "", data["phone"], flags=re.IGNORECASE)
        if data["email"] and "@" not in data["email"]:
            warnings.append({"row": number, "code": "invalid_email", "value": data["email"]})
            data["email"] = ""
        link = data["link"]
        if link and not re.match(r"https?://", link, re.IGNORECASE):
            if "." in link and " " not in link:
                data["link"] = "https://" + link
            else:
                warnings.append({"row": number, "code": "invalid_link", "value": link})
                data["link"] = ""

        data = {key: clip(key, value) for key, value in data.items()}

        dates = {}
        for key in DATE_FIELDS:
            value = get(row, columns, key)
            try:
                dates[key] = parse_date(value)
            except ValueError:
                dates[key] = None
                warnings.append({"row": number, "code": "invalid_date", "value": as_text(value)})

        identity = company_key(data["company"])
        if identity in seen:
            duplicates += 1  # the same company twice in one file, only the first row counts
            continue
        seen.add(identity)

        match = by_company.get(identity)
        if match:
            if not update_existing:
                skipped += 1  # mode "new only": existing companies are left alone
                continue
            # Update: every filled cell of the file replaces the stored value. Empty cells and
            # columns that are not in the file never delete anything. Attachments stay as they are.
            values = {key: value for key, value in data.items() if value}
            values.pop("company")  # same company by definition, keep the spelling that is stored
            values.update({key: value for key, value in dates.items() if value is not None})
            for key, labels in (("status", STATUS_LABELS), ("channel", CHANNEL_LABELS), ("work_mode", WORK_MODE_LABELS)):
                label = labels.get(as_text(get(row, columns, key)).lower())
                if label:
                    values[key] = label
            favorite = as_text(get(row, columns, "favorite")).lower()
            if favorite:
                values["is_favorite"] = favorite in FAVORITE_VALUES

            changed = [key for key, value in values.items() if getattr(match, key) != value]
            if changed:
                to_update.append((match, {key: values[key] for key in changed}))
                if len(update_preview) < PREVIEW_ROWS:
                    update_preview.append({"company": match.company, "fields": changed})
            else:
                unchanged += 1
            continue

        to_create.append(Application(
            **data,
            **dates,
            status=STATUS_LABELS.get(as_text(get(row, columns, "status")).lower(), default_status),
            channel=CHANNEL_LABELS.get(as_text(get(row, columns, "channel")).lower(), ""),
            work_mode=WORK_MODE_LABELS.get(as_text(get(row, columns, "work_mode")).lower(), ""),
            is_favorite=as_text(get(row, columns, "favorite")).lower() in FAVORITE_VALUES,
        ))
        if len(preview) < PREVIEW_ROWS:
            preview.append({"company": data["company"], "position": data["position"], "city": data["city"]})

    if not dry_run:
        with transaction.atomic():
            if to_create:
                Application.objects.bulk_create(to_create)
            for application, changes in to_update:
                for key, value in changes.items():
                    setattr(application, key, value)
                application.save()  # a changed address also resets the stored map coordinates

    return {
        "dry_run": dry_run,
        "created": len(to_create),
        "updated": len(to_update),
        "unchanged": unchanged,
        "skipped": skipped,
        "duplicates": duplicates,
        "errors": errors,
        "warnings": warnings,
        "columns": list(columns),
        "ignored_columns": ignored,
        "preview": preview,
        "update_preview": update_preview,
    }